from __future__ import annotations

import argparse
import csv
import hashlib
import json
from io import BytesIO
from pathlib import Path

from .normalize import canonicalize_row, is_riverview, norm_header
from .workbook import PERSONNEL_FIELDS


def _header_row(rows):
    for number, row in rows:
        normalized = {norm_header(v) for v in row if v not in (None, "")}
        if PERSONNEL_FIELDS.issubset(normalized):
            return number, list(row)
    raise RuntimeError("S-275 personnel header not found")


def iter_workbook(path: Path):
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [(i, tuple(row)) for i, row in enumerate(ws.iter_rows(values_only=True), 1)]
            try:
                header_no, headers = _header_row(rows[:25])
            except RuntimeError:
                continue
            for number, row in rows:
                if number > header_no:
                    yield ws.title, number, headers, list(row)
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path, on_demand=True)
        try:
            for name in wb.sheet_names():
                ws = wb.sheet_by_name(name)
                rows = [(i + 1, ws.row_values(i)) for i in range(ws.nrows)]
                try:
                    header_no, headers = _header_row(rows[:25])
                except RuntimeError:
                    continue
                for number, row in rows:
                    if number > header_no:
                        yield name, number, headers, row
        finally:
            wb.release_resources()
    else:
        raise RuntimeError(f"unsupported workbook: {path}")


def normalize_source(source: dict):
    path = Path(source["local_path"])
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != source["sha256"]:
        raise RuntimeError(f"{source['school_year']}: raw file hash does not match manifest")
    for sheet, number, headers, values in iter_workbook(path):
        row = canonicalize_row(
            school_year=source["school_year"], source_sha256=digest,
            source_sheet=sheet, source_row=number, headers=headers, row=values,
        )
        if is_riverview(row):
            yield row


def summarize(rows):
    by_year = {}
    for row in rows:
        item = by_year.setdefault(row["school_year"], {"rows": 0, "certificated_fte": 0.0, "classified_fte": 0.0, "base_salary": 0.0, "total_salary": 0.0})
        item["rows"] += 1
        for field in ("certificated_fte", "classified_fte", "base_salary", "total_salary"):
            item[field] += row[field] or 0.0
    return by_year


def main(argv=None):
    ap = argparse.ArgumentParser(description="Normalize validated S-275 files and extract RSD407")
    ap.add_argument("--manifest", default="artifacts/manifests/collection.json")
    ap.add_argument("--output", default="artifacts/normalized/rsd407.csv")
    ap.add_argument("--summary", default="artifacts/normalized/rsd407-summary.json")
    args = ap.parse_args(argv)
    manifest = json.loads(Path(args.manifest).read_text())
    rows = []
    for source in manifest["sources"]:
        rows.extend(normalize_source(source))
    expected = set(source["school_year"] for source in manifest["sources"])
    actual = set(row["school_year"] for row in rows)
    if actual != expected:
        raise RuntimeError(f"normalized year coverage mismatch: expected {sorted(expected)}, got {sorted(actual)}")
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    summary = {"schema_version": 1, "district": "Riverview School District", "years": summarize(rows)}
    summary_path = Path(args.summary)
    summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n")
    print(f"normalized {len(rows)} RSD407 source row(s) across {len(actual)} year(s): {output}")


if __name__ == "__main__":
    main()
