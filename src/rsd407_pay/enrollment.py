from __future__ import annotations
import argparse,json
from datetime import datetime,timezone
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from .enrollment_source import fetch_source,sha256

from .period import YEARS, require_coverage
DISTRICT_CODE="17407"
DISTRICT_NAME="Riverview"
SECTION="K-12 FTE - Includes ALE"

def extract(data:bytes):
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    out=[]
    for year in YEARS:
        if year not in wb.sheetnames: raise RuntimeError(f"missing P-223 worksheet {year}")
        ws=wb[year]
        top=list(next(ws.iter_rows(min_row=2,max_row=2,values_only=True)))
        headers=list(next(ws.iter_rows(min_row=3,max_row=3,values_only=True)))
        starts=[i for i,v in enumerate(top) if v==SECTION]
        if starts!=[2]: raise RuntimeError(f"{year} expected {SECTION!r} at column 3, found indexes {starts}")
        end=next((i for i in range(3,len(top)) if top[i] not in (None,"")),None)
        if end is None: raise RuntimeError(f"{year} could not find end of {SECTION!r}")
        grade_headers=headers[2:end]
        if grade_headers[-1]!="12th": raise RuntimeError(f"{year} K-12 section does not end at 12th: {grade_headers!r}")
        matches=[]
        for row in ws.iter_rows(min_row=4,values_only=True):
            code=str(row[0]).strip() if row[0] is not None else ""
            if code==DISTRICT_CODE: matches.append(row)
        if len(matches)!=1: raise RuntimeError(f"{year} expected one district {DISTRICT_CODE}, found {len(matches)}")
        row=matches[0]
        if str(row[1]).strip()!=DISTRICT_NAME: raise RuntimeError(f"{year} district code {DISTRICT_CODE} name changed: {row[1]!r}")
        values=[]
        for h,v in zip(grade_headers,row[2:end]):
            try: values.append(float(v))
            except (TypeError,ValueError): raise RuntimeError(f"{year} nonnumeric K-12 FTE component {h!r}: {v!r}")
        out.append({"school_year":year,"district_code":DISTRICT_CODE,"district_name":DISTRICT_NAME,
                    "measure":"P-223 final annual-average K-12 FTE including ALE",
                    "component_headers":[str(x) for x in grade_headers],"student_fte":sum(values)})
    return out

def main(argv=None):
    p=argparse.ArgumentParser(); p.add_argument("--output",default="artifacts/normalization/enrollment.json"); p.add_argument("--raw-output",default="artifacts/normalization/raw/p223-final-enrollment.xlsx"); args=p.parse_args(argv)
    source,data=fetch_source()
    rows=extract(data)
    raw=Path(args.raw_output)
    raw.parent.mkdir(parents=True,exist_ok=True)
    raw.write_bytes(data)
    payload={"schema_version":1,"source":{"provider":"Washington OSPI","dataset":"P-223 Final Enrollment Summary",
             "landing_page":source.landing_page,"workbook_url":source.workbook_url,
             "retrieved_at":datetime.now(timezone.utc).isoformat(),"sha256":sha256(data),"size_bytes":len(data),"raw_source_file":str(raw)},
             "district":{"code":DISTRICT_CODE,"name":DISTRICT_NAME},"years":rows}
    path=Path(args.output); path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(payload,indent=2)+"\n")
    for r in rows: print(f"{r['school_year']}: {r['student_fte']:.2f} student FTE ({len(r['component_headers'])} K-12 components)")
if __name__=="__main__": main()

