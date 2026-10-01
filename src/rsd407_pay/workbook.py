from __future__ import annotations
from io import BytesIO
from pathlib import Path

PERSONNEL_FIELDS = {
    "school district",
    "duty title",
    "cert fte",
    "clas fte",
    "base salary",
    "total salary",
}
OPTIONAL_PERSONNEL_FIELDS = {"insurance benefits", "mandatory benefits"}

def _norm(value):
    return " ".join(str(value or "").strip().lower().split())

def _matches(values):
    headers={_norm(v) for v in values if v is not None}
    return PERSONNEL_FIELDS.issubset(headers)

def validate_workbook_bytes(data: bytes, extension: str) -> dict:
    ext=extension.lower().lstrip(".")
    if ext=="xlsx":
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=25, values_only=True):
                if _matches(row):
                    return {"format":"xlsx","sheet":ws.title,"headers":[_norm(v) for v in row if v is not None]}
        raise RuntimeError("workbook does not contain an S-275 personnel header signature")
    if ext=="xls":
        import xlrd
        wb=xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for name in wb.sheet_names():
                ws=wb.sheet_by_name(name)
                for i in range(min(ws.nrows,25)):
                    row=ws.row_values(i)
                    if _matches(row):
                        return {"format":"xls","sheet":name,"headers":[_norm(v) for v in row if v not in (None,"")]}
        finally:
            wb.release_resources()
        raise RuntimeError("workbook does not contain an S-275 personnel header signature")
    raise RuntimeError(f"unsupported workbook format for S-275 validation: {extension}")
