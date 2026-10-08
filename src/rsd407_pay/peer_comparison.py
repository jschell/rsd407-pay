"""Compare selected districts using identical retained S-275/P-223 definitions."""
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from .pipeline import iter_workbook
from .normalize import canonicalize_row, norm_header
from .job_family import load_mapping
from .enrollment import SECTION

DISTRICTS = {
 'Riverview': ['Riverview'], 'Snoqualmie Valley': ['Snoqualmie Valley'],
 'Northshore': ['Northshore'], 'Monroe': ['Monroe'], 'Sultan': ['Sultan'],
 'Lakewood': ['Lakewood'], 'Granite Falls': ['Granite Falls'], 'Tukwila': ['Tukwila'],
 'Orting': ['Orting'], 'Steilacoom Historical': ['Steilacoom Hist.', 'Steilacoom Historical', 'Steilacoom'],
}
COMPARISON_ADDITIONS = {'Deputy/Assist. Supt.': 'district/central administration', 'Elem. Vice Principal': 'principals/APs'}
YEARS = ['2013-14', '2023-24', '2024-25']


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def enrollment(data, years=YEARS):
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    aliases={name: district for district,names in DISTRICTS.items() for name in names}
    result={}
    for year in years:
        ws=wb[year]
        top=list(next(ws.iter_rows(min_row=2,max_row=2,values_only=True)))
        headers=list(next(ws.iter_rows(min_row=3,max_row=3,values_only=True)))
        if [i for i,v in enumerate(top) if v==SECTION] != [2]:
            raise RuntimeError(f'{year}: unexpected P-223 K-12 layout')
        end=next(i for i in range(3,len(top)) if top[i] not in (None,''))
        if headers[end-1]!='12th': raise RuntimeError(f'{year}: K-12 section not complete')
        for row in ws.iter_rows(min_row=4,values_only=True):
            district=aliases.get(str(row[1]).strip())
            if district is None: continue
            key=(year,district)
            if key in result: raise RuntimeError(f'duplicate enrollment district: {key}')
            student_fte=sum(float(v) for v in row[2:end])
            if not math.isfinite(student_fte) or student_fte<=0: raise RuntimeError(f'invalid enrollment: {key}')
            result[key]={'student_fte':student_fte,'district_code':str(row[0]),'source_label':str(row[1])}
        missing=[d for d in DISTRICTS if (year,d) not in result]
        if missing: raise RuntimeError(f'{year}: missing enrollment districts {missing}')
    wb.close()
    return result


def aggregate(rows, enroll, mapping, year):
    out={d:{'district':d,'school_year':year,**enroll[(year,d)],
            'district_fte':0.,'district_compensation':0.,'central_fte':0.,'central_salary':0.,
            'central_compensation':0.,'school_admin_fte':0.,'school_admin_compensation':0.,
            'director_supervisor_fte_excluded':0.,'director_supervisor_compensation_excluded':0.,
            'unmapped_title_fte':0.,'unmapped_title_compensation':0.,'employee_rows':0,'central_employee_rows':0,'zero_fte_rows':0,'salary_below_base_rows':0}
         for d in DISTRICTS}
    aliases={norm_header(name):d for d,names in DISTRICTS.items() for name in names}
    seen=set(); titles={d:{} for d in DISTRICTS}
    for row in rows:
        district=aliases.get(norm_header(row['district_name']))
        if district is None: continue
        key=(row['source_sheet'],row['source_row'])
        if key in seen: raise RuntimeError(f'duplicate source row {key}')
        seen.add(key)
        title=str(row['duty_title'] or '').strip()
        values={k:float(row[k] or 0) for k in ['certificated_fte','classified_fte','base_salary','total_salary','insurance_benefits','mandatory_benefits']}
        if any(not math.isfinite(v) or v<0 for v in values.values()): raise RuntimeError(f'{district}: invalid numeric source row')
        fte=values['certificated_fte']+values['classified_fte']
        comp=values['total_salary']+values['insurance_benefits']+values['mandatory_benefits']
        r=out[district];r['employee_rows']+=1;r['district_fte']+=fte;r['district_compensation']+=comp
        r['zero_fte_rows']+=int(fte==0);r['salary_below_base_rows']+=int(values['total_salary']<values['base_salary'])
        t=titles[district].setdefault(title,{'rows':0,'fte':0.,'salary':0.,'compensation':0.})
        t['rows']+=1;t['fte']+=fte;t['salary']+=values['total_salary'];t['compensation']+=comp
        family=mapping.get(title,'unmapped/review')
        if family=='unmapped/review':
            r['unmapped_title_fte']+=fte;r['unmapped_title_compensation']+=comp
        if family=='district/central administration':
            r['central_employee_rows']+=1;r['central_fte']+=fte;r['central_salary']+=values['total_salary'];r['central_compensation']+=comp
        if family=='principals/APs': r['school_admin_fte']+=fte;r['school_admin_compensation']+=comp
        if title=='Director/Supervisor':
            r['director_supervisor_fte_excluded']+=fte;r['director_supervisor_compensation_excluded']+=comp
    for district,r in out.items():
        if not r['employee_rows'] or r['central_fte']<=0: raise RuntimeError(f'{year}: missing source/central staffing {district}')
        r['central_fte_per_1000_students']=r['central_fte']/r['student_fte']*1000
        r['central_compensation_per_1000_students']=r['central_compensation']/r['student_fte']*1000
        r['central_compensation_per_staff_fte']=r['central_compensation']/r['central_fte']
        r['central_share_district_compensation']=r['central_compensation']/r['district_compensation']
    return list(out.values()),titles


def main(argv=None):
    p=argparse.ArgumentParser();p.add_argument('--output',default='artifacts/peer-comparison');a=p.parse_args(argv)
    collection=json.loads(Path('artifacts/manifests/collection.json').read_text())
    if collection['release_scope']!='final': raise RuntimeError('final source required')
    normalization=json.loads(Path('artifacts/normalization/normalization-source-manifest.json').read_text())
    ep=Path('artifacts/normalization/raw/p223-final-enrollment.xlsx')
    expected=next(x['sha256'] for x in normalization['files'] if x['path']==str(ep))
    if sha(ep)!=expected: raise RuntimeError('enrollment hash mismatch')
    enroll=enrollment(ep.read_bytes());mapping={**load_mapping()['mappings'],**COMPARISON_ADDITIONS};all_rows=[];evidence=[];titles={}
    for year in YEARS:
        source=next(x for x in collection['sources'] if x['school_year']==year)
        path=Path(source['local_path'])
        if sha(path)!=source['sha256']: raise RuntimeError(f'{year}: source hash mismatch')
        def selected():
            for sheet,number,headers,values in iter_workbook(path):
                yield canonicalize_row(school_year=year,source_sha256=source['sha256'],source_sheet=sheet,source_row=number,headers=headers,row=values)
        annual,title_rows=aggregate(selected(),enroll,mapping,year)
        all_rows.extend(annual);titles[year]=title_rows;evidence.append(source)
        print(f'{year}: compared {len(annual)} districts',flush=True)
    out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    report={'schema_version':1,'status':'source_verified_comparison','years':YEARS,'rows':all_rows,
            'duty_title_evidence':titles,'s275_sources':evidence,'normalization_manifest':normalization,
            'mapping_sha256':sha('config/job-family-mapping.json'),'comparison_mapping_additions':COMPARISON_ADDITIONS,'comparison_code_sha256':sha(__file__),
            'limitations':['Strict central administration includes Superintendent, Deputy/Assist. Supt., and Other District Admin.; Director/Supervisor is excluded and reported separately.',
              'Employee-year compensation, not operating expenditure; no causal or efficiency conclusion.',
              'Unfamiliar peer Duty Titles are retained in unmapped/review and reported as FTE/compensation; no functional classification is inferred.',
              'Peer totals have not been independently reconciled with published Personnel Summary or Access controls.',
              '2019-20 is omitted here; comparison uses historical baseline and two most recent accepted years.']}
    (out/'comparison.json').write_text(json.dumps(report,indent=2)+'\n')
    with (out/'comparison.csv').open('w',newline='') as h:
        w=csv.DictWriter(h,fieldnames=list(all_rows[0]));w.writeheader();w.writerows(all_rows)
    latest=sorted([r for r in all_rows if r['school_year']==YEARS[-1]],key=lambda r:r['central_compensation_per_1000_students'])
    lines=['# Central administration peer comparison','', '2024-25; salary plus reported employer insurance and mandatory benefits. Student denominator: OSPI P-223 annual-average K-12 FTE including ALE.','',
           '| District | Student FTE | Central FTE | FTE / 1,000 | Compensation / 1,000 | Compensation / admin FTE | Excluded Director/Supervisor FTE |',
           '| --- | --- | --- | --- | --- | --- | --- |']
    for r in latest:
        lines.append(f"| {r['district']} | {r['student_fte']:,.2f} | {r['central_fte']:.3f} | {r['central_fte_per_1000_students']:.2f} | ${r['central_compensation_per_1000_students']:,.0f} | ${r['central_compensation_per_staff_fte']:,.0f} | {r['director_supervisor_fte_excluded']:.3f} |")
    lines+=['','## Limits','']+['- '+x for x in report['limitations']]
    (out/'comparison.md').write_text('\n'.join(lines)+'\n')
if __name__=='__main__':main()
