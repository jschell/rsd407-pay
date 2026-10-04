from __future__ import annotations
import argparse,json
from datetime import datetime,timezone
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from .enrollment_source import fetch_source,sha256

def inventory(data:bytes):
    wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        sample=[]
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,25),values_only=True):
            vals=[None if v is None else str(v).strip() for v in row]
            if any(v not in (None,"") for v in vals): sample.append(vals)
        sheets.append({"title":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,"first_nonempty_rows":sample})
    return {"sheets":sheets}

def main(argv=None):
    ap=argparse.ArgumentParser(); ap.add_argument("--output",default="artifacts/normalization/p223-inventory.json"); a=ap.parse_args()
    source,data=fetch_source()
    report={"source":{"landing_page":source.landing_page,"workbook_url":source.workbook_url,
      "retrieved_at":datetime.now(timezone.utc).isoformat(),"sha256":sha256(data),"size_bytes":len(data)},
      **inventory(data)}
    out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True); out.write_text(json.dumps(report,indent=2)+"\n")
    print(f"P-223 workbook: {source.workbook_url}")
    print(f"sha256={report['source']['sha256']} bytes={len(data)} sheets={len(report['sheets'])}")
    for s in report["sheets"]:
        print(f"sheet={s['title']!r} rows={s['max_row']} cols={s['max_column']}")
        for i,row in enumerate(s["first_nonempty_rows"][:8],1): print(f"  sample[{i}]={row}")
if __name__=="__main__": main()
