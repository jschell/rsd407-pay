import unittest
from io import BytesIO
from openpyxl import Workbook
from rsd407_pay.enrollment import extract,YEARS,DISTRICT_CODE,DISTRICT_NAME

def workbook():
    wb=Workbook(); wb.remove(wb.active)
    for year in YEARS:
        ws=wb.create_sheet(year)
        ws.append([1,2,3,4,5])
        ws.append([f"FINAL {year}",None,"K-12 FTE - Includes ALE",None,"ALE"])
        ws.append(["CCDDD","District","K","12th","K"])
        ws.append([DISTRICT_CODE,DISTRICT_NAME,100,200,999])
    b=BytesIO(); wb.save(b); return b.getvalue()

class EnrollmentTests(unittest.TestCase):
    def test_extracts_only_k12_section(self):
        rows=extract(workbook())
        self.assertEqual(len(rows),12)
        self.assertEqual(rows[0]["student_fte"],300)
        self.assertEqual(rows[0]["component_headers"],["K","12th"])
    def test_requires_exact_district_name(self):
        data=workbook(); wb=load_workbook(BytesIO(data))
        wb[YEARS[0]]["B4"]="Wrong"
        b=BytesIO(); wb.save(b)
        with self.assertRaisesRegex(RuntimeError,"name changed"): extract(b.getvalue())

from openpyxl import load_workbook
if __name__=="__main__": unittest.main()
