import unittest
from openpyxl import Workbook
from rsd407_pay.overhead_review import extract_sheet, ACTIVITIES


class OverheadSourceTests(unittest.TestCase):
    def sheet(self):
        s = Workbook().active
        s.title = 'Activity'
        s.append([None])
        s.append([None, '2024-25 General Fund Expenditures by Activity'])
        s.append([None]*5 + list(ACTIVITIES))
        s.append([None,'County District Code','District Name','Total Full Enrollment','Total'] + list(ACTIVITIES.values()))
        s.append([None])
        s.append([None,'17407','Riverview',3000,100*len(ACTIVITIES)] + [100]*len(ACTIVITIES))
        return s

    def test_reconciliation_and_missing_district(self):
        s=self.sheet()
        result=extract_sheet(s,'2024-25',{'17407':'Riverview'})
        self.assertEqual(result['17407']['total'],100*len(ACTIVITIES))
        self.assertEqual(result['17407']['source_row'],6)
        with self.assertRaises(ValueError):
            extract_sheet(s,'2024-25',{'17407':'Riverview','17410':'Snoqualmie Valley'})
        s.cell(6,5,1300)
        with self.assertRaises(ValueError):
            extract_sheet(s,'2024-25',{'17407':'Riverview'})

    def test_duplicates_schema_and_nonfinite(self):
        s=self.sheet();s.append([c.value for c in s[6]])
        with self.assertRaises(ValueError):extract_sheet(s,'2024-25',{'17407':'Riverview'})
        s=self.sheet();s.cell(4,6,'Unexpected category')
        with self.assertRaises(ValueError):extract_sheet(s,'2024-25',{'17407':'Riverview'})
        s=self.sheet();s.cell(6,6,float('nan'))
        with self.assertRaises(ValueError):extract_sheet(s,'2024-25',{'17407':'Riverview'})


if __name__ == '__main__':
    unittest.main()
